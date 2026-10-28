"""Generate every figure used in Report/so101_report_EN.tex.

Usage:
    python scripts/make_report_figures.py --cache <download dir>

Inputs, all in this repository or public:
  - the three datasets on Hugging Face (tables + first video file only), downloaded to --cache;
  - Report/Evaluation results.xlsx, the per-position evaluation outcomes;
  - Report/media/, photos and video stills extracted once from the original recordings
    (not published) by scripts/extract_media_stills.py.

Needs: ffmpeg on PATH (or --ffmpeg), matplotlib, numpy, pandas, openpyxl, huggingface_hub.
"""
import argparse
import glob
import shutil
import subprocess
from math import sqrt
from pathlib import Path

import matplotlib

matplotlib.use("Agg")
import matplotlib.image as mpimg
import matplotlib.patches
import matplotlib.pyplot as plt
import numpy as np
import openpyxl
import pandas as pd
from huggingface_hub import snapshot_download

REPORT = Path(__file__).resolve().parent.parent / "Report"
OUT = REPORT / "figures"
MEDIA = REPORT / "media"
DATASETS = {
    "d1": "so101_cube_demo_0_20260926_155239",
    "v2": "so101_cube_v2",
    "v3": "so101_cube_v3",
}

# palette: one accent, one highlight, neutral ink; outcome scale is a single-hue ramp
BLUE, ORANGE, GREY = "#2a78d6", "#eb6834", "#8a8984"
INK, INK2 = "#0b0b0b", "#52514e"
OUTCOME = {"a": "#1d5fae", "b": "#5b9be3", "c": "#a9c9ef", "d": "#e4e3df"}
POINTS = {"a": 3, "b": 2, "c": 1, "d": 0}

plt.rcParams.update({
    "font.family": "sans-serif",
    "font.size": 8,
    "axes.edgecolor": INK2,
    "axes.labelcolor": INK,
    "axes.linewidth": 0.6,
    "axes.spines.top": False,
    "axes.spines.right": False,
    "xtick.color": INK2,
    "ytick.color": INK2,
    "xtick.major.width": 0.6,
    "ytick.major.width": 0.6,
    "savefig.bbox": "tight",
    "savefig.pad_inches": 0.02,
})


def wilson(k, n, z=1.96):
    p = k / n
    c = (p + z * z / (2 * n)) / (1 + z * z / n)
    h = z * sqrt(p * (1 - p) / n + z * z / (4 * n * n)) / (1 + z * z / n)
    return c - h, c + h


def frame(ffmpeg, src, t, dst, vf=None):
    cmd = [ffmpeg, "-loglevel", "error", "-y", "-ss", str(t), "-i", str(src), "-frames:v", "1"]
    if vf:
        cmd += ["-vf", vf]
    subprocess.run(cmd + ["-q:v", "2", str(dst)], check=True)


def save_jpg(fig, name, dpi=250):
    fig.savefig(OUT / name, dpi=dpi, pil_kwargs={"quality": 92})
    plt.close(fig)


def gripper(data_dir):
    files = sorted(glob.glob(str(Path(data_dir) / "data" / "*" / "*.parquet")))
    d = pd.concat([pd.read_parquet(f) for f in files])
    return np.stack(d["action"].values)[:, 5]


def panel(ax, img, label=None, fs=7):
    ax.imshow(img)
    ax.set_xticks([])
    ax.set_yticks([])
    for s in ax.spines.values():
        s.set_visible(False)
    if label:
        ax.text(0.02, 0.97, label, transform=ax.transAxes, ha="left", va="top", fontsize=fs,
                color="white", bbox=dict(boxstyle="round,pad=0.25", fc="black", ec="none", alpha=0.55))


# ---------------------------------------------------------------- evaluation
def load_outcomes():
    ws = openpyxl.load_workbook(REPORT / "Evaluation results.xlsx").active
    models = [ws.cell(1, c).value for c in (2, 3, 4)]
    rows = [[ws.cell(r, c).value for c in (2, 3, 4)] for r in range(2, 17)]
    return {m: [row[i] for row in rows] for i, m in enumerate(models)}


def fig_results(outcomes):
    models = ["v2_024", "v2_100", "v3_050"]
    labels = ["v2_024 · 24 demos", "v2_100 · 100 demos", "v3_050 · 51 demos"]
    loss = [0.050, 0.076, 0.09]  # v3 value approximate (15k of 20k steps)
    n = 15

    fig, (a, b) = plt.subplots(1, 2, figsize=(6.3, 2.0), gridspec_kw={"wspace": 0.45})
    y = np.arange(len(models))[::-1]

    a.barh(y, loss, color=GREY, height=0.5)
    for yi, v, approx in zip(y, loss, [False, False, True]):
        a.text(v + 0.003, yi, ("≈" if approx else "") + (f"{v:.2f}" if approx else f"{v:.3f}"),
               va="center", color=INK, fontsize=7)
    a.set_xlim(0, 0.12)
    a.set_yticks(y, labels)
    a.set_xlabel("final training L1 loss")
    a.set_title("Training converges", loc="left", fontsize=9, fontweight="bold", color=INK)

    for yi, m in zip(y, models):
        k = outcomes[m].count("a")
        lo, hi = wilson(k, n)
        b.plot([lo * 100, hi * 100], [yi, yi], color=BLUE, lw=2, solid_capstyle="round", alpha=0.35)
        b.plot(k / n * 100, yi, "o", color=BLUE, ms=6, mec="white", mew=1)
        b.text(hi * 100 + 2, yi, f"{k}/{n}", va="center", color=INK, fontsize=7)
    b.set_xlim(0, 100)
    b.set_yticks(y, [])
    b.set_xlabel("task success, % (Wilson 95% CI)")
    b.set_title("Task success does not", loc="left", fontsize=9, fontweight="bold", color=INK)
    fig.savefig(OUT / "results.pdf")
    plt.close(fig)


def fig_outcomes(outcomes):
    models = ["v2_024", "v2_100", "v3_050"]
    names = {"a": "a · success", "b": "b · lifted, not deposited", "c": "c · near miss (within ≈1 cm)",
             "d": "d · complete failure"}
    fig, ax = plt.subplots(figsize=(6.3, 1.55))
    y = np.arange(len(models))[::-1]
    for yi, m in zip(y, models):
        left = 0
        for cat in "abcd":
            k = outcomes[m].count(cat)
            if k:
                ax.barh(yi, k, left=left, color=OUTCOME[cat], height=0.6, edgecolor="white", linewidth=1.5)
                ax.text(left + k / 2, yi, str(k), ha="center", va="center", fontsize=7,
                        color="white" if cat in "ab" else INK)
            left += k
        pts = sum(POINTS[o] for o in outcomes[m])
        ax.text(15.4, yi, f"{pts} / 45 pts", va="center", fontsize=7.5, color=INK, fontweight="bold")
    ax.set_yticks(y, models)
    ax.set_xlim(0, 15)
    ax.set_xticks([0, 5, 10, 15])
    ax.set_xlabel("rollouts (15 positions)")
    ax.spines["left"].set_visible(False)
    ax.tick_params(axis="y", length=0)
    handles = [matplotlib.patches.Patch(color=OUTCOME[c], label=names[c]) for c in "abcd"]
    ax.legend(handles=handles, loc="lower center", bbox_to_anchor=(0.5, 1.0), ncol=4, frameon=False, fontsize=7,
              handlelength=1.0, columnspacing=1.2)
    fig.savefig(OUT / "outcomes.pdf")
    plt.close(fig)


# ---------------------------------------------------------------- gripper
def fig_gripper(d1_dir, v2_dir):
    d1, v2 = gripper(d1_dir), gripper(v2_dir)
    fig, axes = plt.subplots(1, 2, figsize=(6.3, 1.8), sharey=True, gridspec_kw={"wspace": 0.12})
    bins = np.linspace(0, 100, 41)
    for ax, x, title, c in [
        (axes[0], d1, "Dataset 1: gripper opened late", ORANGE),
        (axes[1], v2, "Dataset 2: gripper opened early", BLUE),
    ]:
        w = np.ones_like(x) / len(x) * 100
        ax.hist(x, bins=bins, weights=w, color=c, edgecolor="white", linewidth=0.5)
        closed = (x < 5).mean() * 100
        ax.text(0.97, 0.9, f"{closed:.0f}% of frames closed (<5)", transform=ax.transAxes,
                ha="right", va="top", color=INK, fontsize=7)
        ax.set_title(title, loc="left", fontsize=8, fontweight="bold", color=INK)
        ax.set_xlabel("gripper action (0 = closed)")
    axes[0].set_ylabel("% of frames")
    fig.savefig(OUT / "gripper.pdf")
    plt.close(fig)


# ---------------------------------------------------------------- camera views
def cam_frames(ffmpeg, dirs, tmp):
    out = {}
    for tag in ["v2", "v3"]:
        for cam in ["front", "top"]:
            src = Path(dirs[tag]) / "videos" / f"observation.images.{cam}" / "chunk-000" / "file-000.mp4"
            for t in [4, 6, 8]:
                dst = tmp / f"{tag}_{cam}_{t}.png"
                frame(ffmpeg, src, t, dst)
                out[(tag, cam, t)] = mpimg.imread(dst)
    return out


def fig_setup(f, tag, name, top=0.0):
    """Setup photo (arrows mark where the cameras look) next to the two camera views the policy receives."""
    img = mpimg.imread(MEDIA / f"setup{tag[1]}_photo.jpg")
    img = img[int(img.shape[0] * top):, :]  # drop the shelf, keep the overhead-camera arrow
    ratio = img.shape[1] / img.shape[0] / (4 / 3)
    fig, ax = plt.subplots(1, 3, figsize=(6.3, 1.65), gridspec_kw={"wspace": 0.03, "width_ratios": [ratio, 1, 1]})
    panel(ax[0], img, "workspace")
    panel(ax[1], f[(tag, "top", 6)], "top camera")
    panel(ax[2], f[(tag, "front", 6)], "front camera")
    save_jpg(fig, name)


def fig_details():
    fig, ax = plt.subplots(1, 2, figsize=(3.0, 1.15), gridspec_kw={"wspace": 0.03})
    panel(ax[0], mpimg.imread(MEDIA / "setup3_cube.jpg"), "green cube", fs=6)
    panel(ax[1], mpimg.imread(MEDIA / "setup3_tips.jpg"), "red gripper tips", fs=6)
    save_jpg(fig, "setup3_details.jpg")


def fig_bottleneck(f):
    fig, ax = plt.subplots(1, 3, figsize=(6.3, 1.6), gridspec_kw={"wspace": 0.03})
    panel(ax[0], f[("v2", "top", 6)], "top · approach")
    panel(ax[1], f[("v2", "top", 8)], "top · grasp")
    panel(ax[2], f[("v2", "front", 8)], "front · grasp")
    # cube location, read off the frames by hand
    for a, (x, y), txt in [
        (ax[0], (455, 185), "cube visible"),
        (ax[1], (452, 190), "partly occluded"),
        (ax[2], (350, 375), "depth ambiguous"),
    ]:
        a.add_patch(plt.Circle((x, y), 34, fill=False, ec=ORANGE, lw=1.4, ls="--"))
        a.annotate(txt, (x, y + 36), (x, 460), ha="center", va="bottom", fontsize=7, color="white",
                   fontweight="bold", arrowprops=dict(arrowstyle="-", color=ORANGE, lw=1))
    save_jpg(fig, "bottleneck.jpg")


# ---------------------------------------------------------------- rollout recordings
# label of each still in Report/media/<prefix>_<n>.jpg, see extract_media_stills.py
ROLLOUT_LABELS = {
    "rollout_success": ["approach", "descend", "grasp", "lift"],
    "rollout_nearmiss": ["approach", "descend", "close beside cube", "lift empty"],
    "retry": ["1 · miss", "2 · leaves empty", "3 · back to cube", "4 · regrasp", "5 · in bowl"],
}


def strip(rows, name, height):
    ncol = max(len(ROLLOUT_LABELS[prefix]) for prefix, _ in rows)
    fig, ax = plt.subplots(len(rows), ncol, figsize=(6.3, height), squeeze=False,
                           gridspec_kw={"wspace": 0.03, "hspace": 0.04})
    for r, (prefix, ylabel) in enumerate(rows):
        for c, lab in enumerate(ROLLOUT_LABELS[prefix]):
            panel(ax[r, c], mpimg.imread(MEDIA / f"{prefix}_{c + 1}.jpg"), lab, fs=6)
        if ylabel:
            ax[r, 0].set_ylabel(ylabel, fontsize=8, color=INK)
    save_jpg(fig, name)


def stills():
    for name in ["thumb_teleop", "thumb_success", "thumb_failure", "thumb_retry", "workspace"]:
        shutil.copyfile(MEDIA / f"{name}.jpg", OUT / f"{name}.jpg")


def main():
    p = argparse.ArgumentParser()
    p.add_argument("--cache", required=True, type=Path)
    p.add_argument("--ffmpeg", default="ffmpeg")
    a = p.parse_args()

    OUT.mkdir(parents=True, exist_ok=True)
    tmp = a.cache / "frames"
    tmp.mkdir(parents=True, exist_ok=True)
    dirs = {
        k: snapshot_download(f"Fabyo04/{r}", repo_type="dataset", local_dir=a.cache / r,
                             allow_patterns=["data/*", "meta/*", "videos/*/chunk-000/file-000.mp4"])
        for k, r in DATASETS.items()
    }

    outcomes = load_outcomes()
    fig_results(outcomes)
    fig_outcomes(outcomes)
    fig_gripper(dirs["d1"], dirs["v2"])
    f = cam_frames(a.ffmpeg, dirs, tmp)
    fig_setup(f, "v2", "setup2.jpg", top=0.05)
    fig_setup(f, "v3", "setup3.jpg", top=0.15)
    fig_details()
    fig_bottleneck(f)
    strip([("rollout_success", "success"), ("rollout_nearmiss", "near miss")], "rollouts.jpg", 1.85)
    strip([("retry", None)], "retry.jpg", 0.95)
    stills()


if __name__ == "__main__":
    main()
